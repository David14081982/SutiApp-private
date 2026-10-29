"""Read-only Excel reconciliation. Private payload never enters a public artifact."""
import argparse, calendar, collections, datetime as dt, hashlib, json, pathlib, re, unicodedata
import openpyxl

def norm(value):
    return ' '.join(str(value or '').strip().casefold().split())

def iso(value):
    return value.isoformat()[:10] if isinstance(value, (dt.datetime, dt.date)) else None

def serial(value):
    return value.isoformat() if isinstance(value, (dt.datetime, dt.date)) else value

def identity(*parts):
    return hashlib.sha256(json.dumps(parts, ensure_ascii=False, separators=(',', ':')).encode()).hexdigest()

def extract(source):
    raw = source.read_bytes()
    formulas = openpyxl.load_workbook(source, data_only=False)
    values = openpyxl.load_workbook(source, data_only=True)
    assert formulas.sheetnames == ['Cartera Inicial', 'Pagos', 'Retiro', 'comisión']
    sha = hashlib.sha256(raw).hexdigest()
    evidence, rows, people, contracts, issues = [], [], {}, [], []
    for sheet in formulas:
        for cells in sheet:
            populated = [c for c in cells if c.value is not None]
            if not populated:
                continue
            evidence.append({'sheet':sheet.title,'row':cells[0].row,'cells':{c.column_letter:{'value':serial(values[sheet.title][c.coordinate].value),'formula':c.value if c.data_type=='f' else None,'raw':serial(c.value),'format':c.number_format,'comment':c.comment.text if c.comment else None} for c in populated}})
        master = sheet.title == 'Cartera Inicial'
        for n, cells in enumerate(values[sheet.title].values, 1):
            c = list(cells)
            name, document = (c[1],c[2]) if master else (c[2],c[3])
            if not isinstance(name, str) or norm(name) == 'inversionista' or len(name.strip()) < 6:
                continue
            if not any(isinstance(v,(dt.date,dt.datetime)) for v in c[:9]) and not (sheet.title=='Pagos' and isinstance(c[10],(int,float))):
                continue
            person = norm(name)
            key = identity('investor', person)
            people.setdefault(key,{'source_key':key,'name':name.strip(),'aliases':set(),'rows':[]})
            people[key]['aliases'].add(name)
            people[key]['rows'].append(f'{sheet.title}!{n}')
            offset = 0 if master else 1
            row = {'sheet':sheet.title,'row':n,'investor_key':key,'name':name,'document':document,'status':c[3+offset], 'start':iso(c[4+offset]),'end':iso(c[5+offset]),'months':c[6+offset], 'principal':c[11] if master else c[8], 'current_capital':c[13] if master else None,'contributed':c[7] if master else None,'withdrawal':c[12] if master else None,'rate':c[14] if master else c[9], 'monthly':c[15] if master else c[10], 'method':c[16] if master else c[11], 'day':c[17] if master else c[12], 'folder':str(c[0] if master else c[1]),'notes':c[13] if sheet.title=='Pagos' else None}
            row['source_key'] = identity('row',sha,sheet.title,n)
            row['kind'] = 'contract' if re.match(r'^contrato\s+\d+',norm(document)) else 'amendment' if norm(document).startswith('adendum') else 'other'
            rows.append(row)
    for row in rows:
        if row['sheet'] != 'Cartera Inicial' or row['kind'] != 'contract':
            continue
        key = identity('contract',row['investor_key'],norm(row['document']),row['start'])
        existing = next((c for c in contracts if c['source_key']==key),None)
        if existing:
            issues.append({'kind':'duplicate_master','source_key':row['source_key'],'contract_key':key,'detail':'Same investor/document/start repeats in master; review before posting.'})
            existing['review'].append('duplicate_master')
            continue
        contract = {**row,'source_key':key,'review':[], 'references':[f"{row['sheet']}!{row['row']}"]}
        if not row['start'] or not row['end'] or row['end']<row['start']:
            contract['review'].append('invalid_dates')
        if not isinstance(row['rate'],(int,float)) or not 0<=row['rate']<=1:
            contract['review'].append('missing_rate')
        if not isinstance(row['principal'],(int,float)) or row['principal']<=0:
            contract['review'].append('missing_principal')
        if row['status']=='Vigente' and row['end'] and row['end']<'2026-10-29':
            contract['review'].append('active_after_maturity')
        contracts.append(contract)
    # Cross-sheet evidence never creates a payment or a second contract.
    for row in rows:
        if row['sheet']=='Cartera Inicial' and row['kind']=='contract':
            continue
        candidates = [c for c in contracts if c['investor_key']==row['investor_key'] and norm(c['document'])==norm(row['document'])]
        exact = [c for c in candidates if c['start']==row['start']]
        if len(exact)==1:
            candidates=exact
        row['contract_candidates']=[c['source_key'] for c in candidates]
        if row['kind']=='amendment':
            # Keep amendment as evidence, never as a parallel capital-bearing contract.
            issues.append({'kind':'amendment_relationship','source_key':row['source_key'],'investor_key':row['investor_key'],'detail':'Amendment requires explicit base-contract link and financial effective date; retained without inventing a second contract.'})
        elif len(candidates)==1:
            c=candidates[0]
            c['references'].append(f"{row['sheet']}!{row['row']}")
            for field in ['start','end','months','principal','rate','monthly','status']:
                a,b=c.get(field),row.get(field)
                same=abs(a-b)<0.011 if isinstance(a,(int,float)) and isinstance(b,(int,float)) else a==b
                if not same:
                    c['review'].append('conflict_'+field)
                    issues.append({'kind':'conflicting_'+field,'source_key':row['source_key'],'contract_key':c['source_key'],'detail':{'field':field,'master':a,'other':b,'master_row':c['row'],'other_sheet':row['sheet'],'other_row':row['row']}})
        else:
            issues.append({'kind':'unlinked_reference' if row['kind']!='other' else 'non_investment_or_unclassified','source_key':row['source_key'],'investor_key':row['investor_key'],'detail':'No unique exact contract match. Preserved for review; no financial posting.'})
    for c in contracts:
        # Annual totals and monthly matrices are historical amounts, not proven dated payments.
        c['review']=sorted(set(c['review']))
    investment_keys={r['investor_key'] for r in rows if r['kind'] in ('contract','amendment')}
    people={key:p for key,p in people.items() if key in investment_keys}
    for p in people.values():
        p['aliases']=sorted(p['aliases'])
        master_rows=[r for r in rows if r['sheet']=='Cartera Inicial' and r['investor_key']==p['source_key']]
        unique={identity(r['investor_key'],norm(r['document']),r['start']):r for r in master_rows}
        p['historical_contributed']=round(sum(r['contributed'] for r in unique.values() if isinstance(r['contributed'],(int,float))),2)
    cutoff='2026-10-29'
    for c in contracts:
        if c['status']=='Vigente':
            amendments=[r for r in rows if r['sheet']=='Cartera Inicial' and r['kind']=='amendment' and r['status']=='Vigente' and r['investor_key']==c['investor_key']]
            if amendments:
                c['review'].append('active_amendment_relationship')
            if isinstance(c['current_capital'],(int,float)) and isinstance(c['rate'],(int,float)) and isinstance(c['monthly'],(int,float)) and abs(c['current_capital']*c['rate']-c['monthly'])>.011:
                c['review'].append('monthly_amount_does_not_match_capital_rate')
            if c['start'] and isinstance(c['day'],(int,float)) and dt.date.fromisoformat(c['start']).day!=int(c['day']):
                c['review'].append('payment_day_differs_from_anniversary')
            if c['start'] and isinstance(c['months'],int):
                start=dt.date.fromisoformat(c['start']);month=start.month-1+c['months'];year=start.year+month//12;month=month%12+1
                expected=dt.date(year,month,min(start.day,calendar.monthrange(year,month)[1])).isoformat()
                if c['end']!=expected:c['review'].append('term_date_mismatch')
        c['review']=sorted(set(c['review']))
        c['activate_snapshot']=c['status']=='Vigente' and not c['review'] and c['start']<=cutoff<c['end'] and isinstance(c['current_capital'],(int,float)) and c['current_capital']>0
        c['future_schedule_from']=None
        if c['activate_snapshot']:
            start=dt.date.fromisoformat(c['start'])
            for n in range(c['months']):
                month=start.month-1+n;year=start.year+month//12;month=month%12+1
                a=dt.date(year,month,min(start.day,calendar.monthrange(year,month)[1])).isoformat()
                month=start.month-1+n+1;year=start.year+month//12;month=month%12+1
                b=dt.date(year,month,min(start.day,calendar.monthrange(year,month)[1])).isoformat()
                if b>=cutoff:c['future_schedule_from']=a;break
    summary={'sha256':sha,'snapshot_observed_on':cutoff,'snapshot_date_meaning':'Date of migration observation, not an inferred cash transaction date','active_snapshots':sum(c['activate_snapshot'] for c in contracts),'sheets':{s.title:{'max_rows':s.max_row,'max_columns':s.max_column,'populated_rows':sum(e['sheet']==s.title for e in evidence),'business_rows':sum(r['sheet']==s.title for r in rows)} for s in formulas},'investors':len(people),'contracts':len(contracts),'contracts_with_conflicts':sum(bool(c['review']) for c in contracts),'amendment_references':sum(r['kind']=='amendment' for r in rows),'issues':dict(collections.Counter(i['kind'] for i in issues)),'realized_payments_inferred':0,'folder_is_unique_identity':False,'workbook_modified':False}
    return {'sha256':sha,'filename':source.name,'summary':summary,'investors':list(people.values()),'contracts':contracts,'rows':rows,'evidence':evidence,'issues':issues}

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('source');parser.add_argument('--output',default='.tmp/inversores/import.json');args=parser.parse_args()
    result=extract(pathlib.Path(args.source)); dest=pathlib.Path(args.output);dest.parent.mkdir(parents=True,exist_ok=True);dest.write_text(json.dumps(result,ensure_ascii=False,default=serial),encoding='utf-8')
    print(json.dumps(result['summary'],ensure_ascii=True,indent=2))
